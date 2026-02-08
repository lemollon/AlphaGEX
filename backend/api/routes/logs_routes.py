"""
Comprehensive Logs API Routes

Provides unified access to ALL logging tables in AlphaGEX:
- Trading Decisions (trading_decisions)
- ML Logs (ml_decision_logs, ml_predictions, fortress_ml_outcomes, spx_wheel_ml_outcomes)
- Prophet Predictions (prophet_predictions)
- Psychology Analysis (psychology_analysis, pattern_learning)
- System Logs (autonomous_trader_logs, spx_debug_logs, data_collection_log)
- AI Analysis (ai_analysis_history, ai_predictions, ai_performance, ai_recommendations)
- Wheel Activity (wheel_activity_log)
- GEX Changes (gex_change_log)
"""

import logging
from datetime import datetime, timedelta
from typing import Optional, List, Dict, Any

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
import io
import csv
import json

from database_adapter import get_connection

router = APIRouter(prefix="/api/logs", tags=["Logs"])
logger = logging.getLogger(__name__)


# ============================================================================
# UNIFIED LOG SUMMARY
# ============================================================================

@router.get("/summary")
async def get_all_logs_summary(days: int = 7):
    """
    Get summary of ALL log tables with record counts and latest entries.
    This gives a complete picture of system activity.
    """
    conn = get_connection()
    cursor = conn.cursor()

    try:
        summaries = {}

        # List of all log tables to check
        log_tables = [
            ('trading_decisions', 'Trading Decisions', 'timestamp'),
            ('autonomous_trader_logs', 'Autonomous Trader', 'timestamp'),
            ('ml_decision_logs', 'ML Decision Logs', 'timestamp'),
            ('ml_predictions', 'ML Predictions', 'timestamp'),
            ('prophet_predictions', 'Prophet Predictions', 'created_at'),
            ('fortress_ml_outcomes', 'FORTRESS ML Outcomes', 'trade_date'),
            ('spx_wheel_ml_outcomes', 'SPX Wheel ML', 'trade_date'),
            ('psychology_analysis', 'Psychology Analysis', 'timestamp'),
            ('pattern_learning', 'Pattern Learning', 'last_seen'),
            ('ai_analysis_history', 'AI Analysis', 'timestamp'),
            ('ai_predictions', 'AI Predictions', 'timestamp'),
            ('ai_performance', 'AI Performance', 'date'),
            ('ai_recommendations', 'AI Recommendations', 'timestamp'),
            ('wheel_activity_log', 'Wheel Activity', 'timestamp'),
            ('gex_change_log', 'GEX Changes', 'timestamp'),
            ('spx_debug_logs', 'SPX Debug', 'timestamp'),
            ('data_collection_log', 'Data Collection', 'timestamp'),
            ('options_collection_log', 'Options Collection', 'timestamp'),
        ]

        total_records = 0

        # PERFORMANCE FIX: Batch check which tables exist (1 query instead of 18)
        table_names = [t[0] for t in log_tables]
        cursor.execute("""
            SELECT table_name FROM information_schema.tables
            WHERE table_schema = 'public' AND table_name = ANY(%s)
        """, (table_names,))
        existing_tables = {row[0] for row in cursor.fetchall()}

        # Build lookup for display names and timestamp columns
        table_info = {t[0]: {'display_name': t[1], 'ts_col': t[2]} for t in log_tables}

        for table_name in table_names:
            info = table_info[table_name]
            display_name = info['display_name']
            ts_col = info['ts_col']

            if table_name not in existing_tables:
                summaries[table_name] = {
                    'display_name': display_name,
                    'exists': False,
                    'total_count': 0,
                    'recent_count': 0,
                    'latest_entry': None
                }
                continue

            try:
                # PERFORMANCE FIX: Single query for all stats (was 3 queries per table)
                cursor.execute(f"""
                    SELECT
                        COUNT(*) as total_count,
                        COUNT(*) FILTER (WHERE {ts_col} >= NOW() - INTERVAL '%s days') as recent_count,
                        MAX({ts_col}) as latest_entry
                    FROM {table_name}
                """, (days,))
                row = cursor.fetchone()
                total, recent, latest = row[0], row[1], row[2]

                summaries[table_name] = {
                    'display_name': display_name,
                    'exists': True,
                    'total_count': total,
                    'recent_count': recent,
                    'latest_entry': str(latest) if latest else None
                }
                total_records += total

            except Exception as e:
                logger.warning(f"Error checking {table_name}: {e}")
                summaries[table_name] = {
                    'display_name': display_name,
                    'exists': False,
                    'error': str(e)
                }

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'total_records_all_tables': total_records,
                'days_analyzed': days,
                'tables': summaries,
                'generated_at': datetime.now().isoformat()
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting logs summary: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# ML LOGS (ml_decision_logs, ml_predictions)
# ============================================================================

@router.get("/ml")
async def get_ml_logs(
    limit: int = Query(50, le=500),
    offset: int = 0,
    action: Optional[str] = None,
    symbol: Optional[str] = None
):
    """Get ML decision logs with filtering."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        # Check if table exists
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'ml_decision_logs'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'logs': [], 'total': 0, 'message': 'Table not yet created'}}

        where_clauses = []
        params = []

        if action:
            where_clauses.append("action = %s")
            params.append(action)
        if symbol:
            where_clauses.append("symbol = %s")
            params.append(symbol)

        where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

        # Get total count
        cursor.execute(f"SELECT COUNT(*) FROM ml_decision_logs {where_sql}", params)
        total = cursor.fetchone()[0]

        # Get logs
        cursor.execute(f"""
            SELECT id, timestamp, action, symbol, details, ml_score,
                   recommendation, reasoning, trade_id, backtest_id
            FROM ml_decision_logs
            {where_sql}
            ORDER BY timestamp DESC
            LIMIT %s OFFSET %s
        """, params + [limit, offset])

        columns = ['id', 'timestamp', 'action', 'symbol', 'details', 'ml_score',
                   'recommendation', 'reasoning', 'trade_id', 'backtest_id']
        logs = [dict(zip(columns, row)) for row in cursor.fetchall()]

        # Convert timestamps
        for log in logs:
            if log['timestamp']:
                log['timestamp'] = str(log['timestamp'])
            if log['details'] and isinstance(log['details'], str):
                try:
                    log['details'] = json.loads(log['details'])
                except:
                    pass

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'logs': logs,
                'total': total,
                'limit': limit,
                'offset': offset
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting ML logs: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# PROPHET PREDICTIONS
# ============================================================================

@router.get("/prophet")
async def get_prophet_predictions(
    limit: int = Query(50, le=500),
    bot_name: Optional[str] = None,
    days: int = 30
):
    """Get Prophet ML predictions with outcomes."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'prophet_predictions'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'predictions': [], 'total': 0, 'message': 'Table not yet created'}}

        where_clauses = ["created_at >= NOW() - INTERVAL '%s days'"]
        params = [days]

        if bot_name:
            where_clauses.append("bot_name = %s")
            params.append(bot_name)

        where_sql = f"WHERE {' AND '.join(where_clauses)}"

        cursor.execute(f"""
            SELECT id, trade_date, bot_name, advice, win_probability, confidence,
                   suggested_risk_pct, reasoning, model_version, top_factors,
                   actual_outcome, actual_pnl, created_at
            FROM prophet_predictions
            {where_sql}
            ORDER BY created_at DESC
            LIMIT %s
        """, params + [limit])

        columns = ['id', 'trade_date', 'bot_name', 'advice', 'win_probability', 'confidence',
                   'suggested_risk_pct', 'reasoning', 'model_version', 'top_factors',
                   'actual_outcome', 'actual_pnl', 'created_at']
        predictions = [dict(zip(columns, row)) for row in cursor.fetchall()]

        for pred in predictions:
            for key in ['trade_date', 'created_at']:
                if pred.get(key):
                    pred[key] = str(pred[key])

        # Get accuracy stats
        cursor.execute("""
            SELECT
                COUNT(*) as total,
                SUM(CASE WHEN actual_outcome = 'WIN' THEN 1 ELSE 0 END) as wins,
                AVG(confidence) as avg_confidence
            FROM prophet_predictions
            WHERE actual_outcome IS NOT NULL
        """)
        stats = cursor.fetchone()

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'predictions': predictions,
                'stats': {
                    'total_with_outcome': stats[0] or 0,
                    'wins': stats[1] or 0,
                    'accuracy': (stats[1] / stats[0] * 100) if stats[0] else 0,
                    'avg_confidence': float(stats[2]) if stats[2] else 0
                }
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting Prophet predictions: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# PSYCHOLOGY ANALYSIS
# ============================================================================

@router.get("/psychology")
async def get_psychology_logs(
    limit: int = Query(50, le=500),
    regime_type: Optional[str] = None,
    days: int = 30
):
    """Get psychology analysis and trap detection logs."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'psychology_analysis'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'analyses': [], 'total': 0, 'message': 'Table not yet created'}}

        where_clauses = ["timestamp >= NOW() - INTERVAL '%s days'"]
        params = [days]

        if regime_type:
            where_clauses.append("regime_type = %s")
            params.append(regime_type)

        where_sql = f"WHERE {' AND '.join(where_clauses)}"

        cursor.execute(f"""
            SELECT id, timestamp, symbol, regime_type, confidence,
                   psychology_trap, trap_probability, market_context,
                   recommended_action, notes
            FROM psychology_analysis
            {where_sql}
            ORDER BY timestamp DESC
            LIMIT %s
        """, params + [limit])

        columns = ['id', 'timestamp', 'symbol', 'regime_type', 'confidence',
                   'psychology_trap', 'trap_probability', 'market_context',
                   'recommended_action', 'notes']
        analyses = [dict(zip(columns, row)) for row in cursor.fetchall()]

        for item in analyses:
            if item.get('timestamp'):
                item['timestamp'] = str(item['timestamp'])

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'analyses': analyses,
                'total': len(analyses)
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting psychology logs: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# AUTONOMOUS TRADER LOGS
# ============================================================================

@router.get("/autonomous")
async def get_autonomous_trader_logs(
    limit: int = Query(50, le=500),
    log_type: Optional[str] = None,
    session_id: Optional[str] = None,
    days: int = 7
):
    """Get autonomous trader scan logs with AI thought process."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'autonomous_trader_logs'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'logs': [], 'total': 0, 'message': 'Table not yet created'}}

        where_clauses = ["timestamp >= NOW() - INTERVAL '%s days'"]
        params = [days]

        if log_type:
            where_clauses.append("log_type = %s")
            params.append(log_type)
        if session_id:
            where_clauses.append("session_id = %s")
            params.append(session_id)

        where_sql = f"WHERE {' AND '.join(where_clauses)}"

        cursor.execute(f"""
            SELECT id, timestamp, log_type, symbol, pattern_detected,
                   confidence_score, ai_thought_process, reasoning_summary,
                   scan_cycle, session_id, additional_data
            FROM autonomous_trader_logs
            {where_sql}
            ORDER BY timestamp DESC
            LIMIT %s
        """, params + [limit])

        columns = ['id', 'timestamp', 'log_type', 'symbol', 'pattern_detected',
                   'confidence_score', 'ai_thought_process', 'reasoning_summary',
                   'scan_cycle', 'session_id', 'additional_data']
        logs = [dict(zip(columns, row)) for row in cursor.fetchall()]

        for log in logs:
            if log.get('timestamp'):
                log['timestamp'] = str(log['timestamp'])

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'logs': logs,
                'total': len(logs)
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting autonomous logs: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# AI RECOMMENDATIONS
# ============================================================================

@router.get("/ai-recommendations")
async def get_ai_recommendations(
    limit: int = Query(50, le=500),
    symbol: Optional[str] = None,
    days: int = 30
):
    """Get AI trade recommendations with outcomes."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'ai_recommendations'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'recommendations': [], 'total': 0, 'message': 'Table not yet created'}}

        where_clauses = ["timestamp >= NOW() - INTERVAL '%s days'"]
        params = [days]

        if symbol:
            where_clauses.append("symbol = %s")
            params.append(symbol)

        where_sql = f"WHERE {' AND '.join(where_clauses)}"

        cursor.execute(f"""
            SELECT id, timestamp, symbol, recommendation_type, action,
                   strike, expiration, confidence, reasoning,
                   expected_return, outcome, actual_pnl
            FROM ai_recommendations
            {where_sql}
            ORDER BY timestamp DESC
            LIMIT %s
        """, params + [limit])

        columns = ['id', 'timestamp', 'symbol', 'recommendation_type', 'action',
                   'strike', 'expiration', 'confidence', 'reasoning',
                   'expected_return', 'outcome', 'actual_pnl']
        recs = [dict(zip(columns, row)) for row in cursor.fetchall()]

        for rec in recs:
            if rec.get('timestamp'):
                rec['timestamp'] = str(rec['timestamp'])
            if rec.get('expiration'):
                rec['expiration'] = str(rec['expiration'])

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'recommendations': recs,
                'total': len(recs)
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting AI recommendations: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# WHEEL ACTIVITY LOG
# ============================================================================

@router.get("/wheel-activity")
async def get_wheel_activity(
    limit: int = Query(50, le=500),
    action: Optional[str] = None,
    days: int = 30
):
    """Get wheel strategy activity logs."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'wheel_activity_log'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'activities': [], 'total': 0, 'message': 'Table not yet created'}}

        where_clauses = ["timestamp >= NOW() - INTERVAL '%s days'"]
        params = [days]

        if action:
            where_clauses.append("action = %s")
            params.append(action)

        where_sql = f"WHERE {' AND '.join(where_clauses)}"

        cursor.execute(f"""
            SELECT id, timestamp, action, description, premium_impact,
                   pnl_impact, underlying_price, option_price, details
            FROM wheel_activity_log
            {where_sql}
            ORDER BY timestamp DESC
            LIMIT %s
        """, params + [limit])

        columns = ['id', 'timestamp', 'action', 'description', 'premium_impact',
                   'pnl_impact', 'underlying_price', 'option_price', 'details']
        activities = [dict(zip(columns, row)) for row in cursor.fetchall()]

        for act in activities:
            if act.get('timestamp'):
                act['timestamp'] = str(act['timestamp'])

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'activities': activities,
                'total': len(activities)
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting wheel activity: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# GEX CHANGE LOG
# ============================================================================

@router.get("/gex-changes")
async def get_gex_changes(
    limit: int = Query(100, le=500),
    symbol: Optional[str] = None,
    change_type: Optional[str] = None,
    days: int = 7
):
    """Get GEX change and direction shift logs."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'gex_change_log'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'changes': [], 'total': 0, 'message': 'Table not yet created'}}

        where_clauses = ["timestamp >= NOW() - INTERVAL '%s days'"]
        params = [days]

        if symbol:
            where_clauses.append("symbol = %s")
            params.append(symbol)
        if change_type:
            where_clauses.append("change_type = %s")
            params.append(change_type)

        where_sql = f"WHERE {' AND '.join(where_clauses)}"

        cursor.execute(f"""
            SELECT id, timestamp, symbol, change_type, previous_value,
                   new_value, change_pct, velocity_trend, direction_change
            FROM gex_change_log
            {where_sql}
            ORDER BY timestamp DESC
            LIMIT %s
        """, params + [limit])

        columns = ['id', 'timestamp', 'symbol', 'change_type', 'previous_value',
                   'new_value', 'change_pct', 'velocity_trend', 'direction_change']
        changes = [dict(zip(columns, row)) for row in cursor.fetchall()]

        for change in changes:
            if change.get('timestamp'):
                change['timestamp'] = str(change['timestamp'])

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'changes': changes,
                'total': len(changes)
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting GEX changes: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# FORTRESS ML OUTCOMES
# ============================================================================

@router.get("/fortress-ml")
async def get_fortress_ml_outcomes(
    limit: int = Query(50, le=500),
    days: int = 30
):
    """Get FORTRESS ML model outcomes for learning analysis."""
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'fortress_ml_outcomes'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'outcomes': [], 'total': 0, 'message': 'Table not yet created'}}

        cursor.execute("""
            SELECT id, trade_date, vix, net_gex, gex_regime, day_of_week,
                   predicted_advice, win_probability, confidence,
                   actual_outcome, is_win, net_pnl, created_at
            FROM fortress_ml_outcomes
            WHERE trade_date >= NOW() - INTERVAL '%s days'
            ORDER BY trade_date DESC
            LIMIT %s
        """, (days, limit))

        columns = ['id', 'trade_date', 'vix', 'net_gex', 'gex_regime', 'day_of_week',
                   'predicted_advice', 'win_probability', 'confidence',
                   'actual_outcome', 'is_win', 'net_pnl', 'created_at']
        outcomes = [dict(zip(columns, row)) for row in cursor.fetchall()]

        for outcome in outcomes:
            if outcome.get('trade_date'):
                outcome['trade_date'] = str(outcome['trade_date'])
            if outcome.get('created_at'):
                outcome['created_at'] = str(outcome['created_at'])

        # Get accuracy stats
        cursor.execute("""
            SELECT
                COUNT(*) as total,
                SUM(CASE WHEN is_win THEN 1 ELSE 0 END) as wins,
                AVG(net_pnl) as avg_pnl
            FROM fortress_ml_outcomes
            WHERE actual_outcome IS NOT NULL
        """)
        stats = cursor.fetchone()

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'outcomes': outcomes,
                'stats': {
                    'total': stats[0] or 0,
                    'wins': stats[1] or 0,
                    'win_rate': (stats[1] / stats[0] * 100) if stats[0] else 0,
                    'avg_pnl': float(stats[2]) if stats[2] else 0
                }
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting FORTRESS ML outcomes: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================================
# EXPORT ALL LOGS
# ============================================================================

@router.get("/export/{log_type}")
async def export_logs(
    log_type: str,
    format: str = Query("csv", regex="^(csv|json)$"),
    days: int = 30
):
    """
    Export logs to CSV or JSON.

    log_type options:
    - trading_decisions
    - ml_logs
    - prophet
    - psychology
    - autonomous
    - ai_recommendations
    - wheel_activity
    - gex_changes
    - ares_ml
    """
    # Get data based on log type
    if log_type == "trading_decisions":
        from trading.decision_logger import export_decisions_json, export_decisions_csv
        if format == "csv":
            content = export_decisions_csv()
            return StreamingResponse(
                io.StringIO(content),
                media_type="text/csv",
                headers={"Content-Disposition": f"attachment; filename=trading_decisions_{datetime.now().strftime('%Y%m%d')}.csv"}
            )
        else:
            data = export_decisions_json(limit=10000)
            return data

    # For other log types, use the respective endpoints
    # This is a simplified version - could be expanded
    raise HTTPException(status_code=400, detail=f"Export not yet implemented for {log_type}")


# ============================================================================
# BOT DECISION LOGS - COMPREHENSIVE UNIFIED LOGGING
# ============================================================================

@router.get("/bot-decisions")
async def get_bot_decisions(
    bot: Optional[str] = None,
    decision_type: Optional[str] = None,
    session_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    outcome: Optional[str] = None,  # 'profit', 'loss', 'pending'
    limit: int = Query(50, le=500),
    offset: int = 0,
    search: Optional[str] = None  # Full-text search
):
    """
    Get comprehensive bot decisions with full filtering.

    This is THE main endpoint for the unified logging system.
    Returns all decision details including Claude prompts, execution timeline, etc.
    """
    conn = get_connection()
    cursor = conn.cursor()

    try:
        # Check if table exists
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'bot_decision_logs'
            )
        """)
        if not cursor.fetchone()[0]:
            return {
                'success': True,
                'data': {
                    'decisions': [],
                    'total': 0,
                    'message': 'bot_decision_logs table not yet created. Run database migration.'
                }
            }

        where_clauses = []
        params = []

        if bot:
            where_clauses.append("bot_name = %s")
            params.append(bot.upper())

        if decision_type:
            where_clauses.append("decision_type = %s")
            params.append(decision_type.upper())

        if session_id:
            where_clauses.append("session_id = %s")
            params.append(session_id)

        if start_date:
            where_clauses.append("timestamp >= %s")
            params.append(start_date)

        if end_date:
            where_clauses.append("timestamp <= %s")
            params.append(end_date + " 23:59:59")

        if outcome:
            if outcome == 'profit':
                where_clauses.append("actual_pnl > 0")
            elif outcome == 'loss':
                where_clauses.append("actual_pnl < 0")
            elif outcome == 'pending':
                where_clauses.append("actual_pnl IS NULL")

        if search:
            where_clauses.append("""
                (entry_reasoning ILIKE %s OR
                 strike_reasoning ILIKE %s OR
                 claude_response ILIKE %s OR
                 outcome_notes ILIKE %s)
            """)
            search_pattern = f"%{search}%"
            params.extend([search_pattern, search_pattern, search_pattern, search_pattern])

        where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

        # Get total count
        cursor.execute(f"SELECT COUNT(*) FROM bot_decision_logs {where_sql}", params)
        total = cursor.fetchone()[0]

        # Get decisions with all fields
        cursor.execute(f"""
            SELECT
                decision_id, bot_name, session_id, scan_cycle, decision_sequence,
                timestamp, decision_type, action, symbol, strategy,
                strike, expiration, option_type, contracts,
                spot_price, vix, net_gex, gex_regime, flip_point, call_wall, put_wall, trend,
                claude_prompt, claude_response, claude_model, claude_tokens_used, claude_response_time_ms,
                langchain_chain, ai_confidence, ai_warnings,
                entry_reasoning, strike_reasoning, size_reasoning, exit_reasoning,
                alternatives_considered, other_strategies_considered,
                psychology_pattern, liberation_setup, false_floor_detected, forward_magnets,
                kelly_pct, position_size_dollars, max_risk_dollars,
                backtest_win_rate, backtest_expectancy, backtest_sharpe,
                risk_checks_performed, passed_all_checks, blocked_reason,
                order_submitted_at, order_filled_at, broker_order_id,
                expected_fill_price, actual_fill_price, slippage_pct, broker_status, execution_notes,
                actual_pnl, exit_triggered_by, exit_timestamp, exit_price, exit_slippage_pct,
                outcome_correct, outcome_notes,
                api_calls_made, errors_encountered, processing_time_ms,
                created_at
            FROM bot_decision_logs
            {where_sql}
            ORDER BY timestamp DESC
            LIMIT %s OFFSET %s
        """, params + [limit, offset])

        columns = [desc[0] for desc in cursor.description]
        decisions = []

        for row in cursor.fetchall():
            d = dict(zip(columns, row))

            # Convert datetime objects to strings
            for key in ['timestamp', 'expiration', 'order_submitted_at', 'order_filled_at',
                       'exit_timestamp', 'created_at']:
                if d.get(key):
                    d[key] = str(d[key])

            # Parse JSON fields
            for key in ['ai_warnings', 'alternatives_considered', 'other_strategies_considered',
                       'forward_magnets', 'risk_checks_performed', 'api_calls_made', 'errors_encountered']:
                if d.get(key) and isinstance(d[key], str):
                    try:
                        d[key] = json.loads(d[key])
                    except:
                        pass

            decisions.append(d)

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'decisions': decisions,
                'total': total,
                'limit': limit,
                'offset': offset,
                'filters_applied': {
                    'bot': bot,
                    'decision_type': decision_type,
                    'session_id': session_id,
                    'start_date': start_date,
                    'end_date': end_date,
                    'outcome': outcome,
                    'search': search
                }
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting bot decisions: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/bot-decisions/{bot_name}")
async def get_bot_specific_decisions(
    bot_name: str,
    limit: int = Query(50, le=500),
    decision_type: Optional[str] = None,
    days: int = 30
):
    """
    Get decisions for a specific bot.

    Shortcut endpoint for per-bot log pages.
    """
    return await get_bot_decisions(
        bot=bot_name,
        decision_type=decision_type,
        start_date=(datetime.now() - timedelta(days=days)).strftime('%Y-%m-%d'),
        limit=limit
    )


@router.get("/bot-decisions/session/{session_id}")
async def get_session_decisions(session_id: str):
    """
    Get all decisions from a specific session.

    Sessions group decisions by date + time period (AM/PM).
    Example session_id: "2024-12-12-AM"
    """
    return await get_bot_decisions(session_id=session_id, limit=100)


@router.get("/bot-decisions/decision/{decision_id}")
async def get_decision_detail(decision_id: str):
    """
    Get full details for a single decision.

    Returns EVERYTHING including full Claude prompt/response.
    """
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT * FROM bot_decision_logs WHERE decision_id = %s
        """, (decision_id,))

        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail=f"Decision {decision_id} not found")

        columns = [desc[0] for desc in cursor.description]
        decision = dict(zip(columns, row))

        # Convert datetime objects
        for key in decision:
            if isinstance(decision[key], datetime):
                decision[key] = decision[key].isoformat()

        # Parse JSON fields
        json_fields = ['ai_warnings', 'alternatives_considered', 'other_strategies_considered',
                      'forward_magnets', 'risk_checks_performed', 'api_calls_made',
                      'errors_encountered', 'full_decision', 'rejection_reasons']
        for key in json_fields:
            if decision.get(key) and isinstance(decision[key], str):
                try:
                    decision[key] = json.loads(decision[key])
                except:
                    pass

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': decision
        }

    except HTTPException:
        raise
    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting decision detail: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/bot-decisions/stats")
async def get_bot_decision_stats(
    bot: Optional[str] = None,
    days: int = 30
):
    """
    Get aggregated statistics for bot decisions.

    Returns:
    - Total decisions by type
    - Win/loss rate
    - Average P&L
    - Decisions per session
    - Error rate
    """
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'bot_decision_logs'
            )
        """)
        if not cursor.fetchone()[0]:
            return {'success': True, 'data': {'message': 'Table not created yet'}}

        bot_filter = "AND bot_name = %s" if bot else ""
        params = [days]
        if bot:
            params.append(bot.upper())

        # Overall stats
        cursor.execute(f"""
            SELECT
                COUNT(*) as total_decisions,
                COUNT(DISTINCT session_id) as total_sessions,
                COUNT(CASE WHEN decision_type = 'ENTRY' THEN 1 END) as entry_decisions,
                COUNT(CASE WHEN decision_type = 'EXIT' THEN 1 END) as exit_decisions,
                COUNT(CASE WHEN decision_type = 'SKIP' THEN 1 END) as skip_decisions,
                COUNT(CASE WHEN actual_pnl > 0 THEN 1 END) as profitable_trades,
                COUNT(CASE WHEN actual_pnl < 0 THEN 1 END) as losing_trades,
                COUNT(CASE WHEN actual_pnl IS NOT NULL THEN 1 END) as closed_trades,
                AVG(actual_pnl) FILTER (WHERE actual_pnl IS NOT NULL) as avg_pnl,
                SUM(actual_pnl) FILTER (WHERE actual_pnl IS NOT NULL) as total_pnl,
                AVG(slippage_pct) FILTER (WHERE slippage_pct IS NOT NULL) as avg_slippage,
                AVG(claude_response_time_ms) FILTER (WHERE claude_response_time_ms > 0) as avg_claude_time,
                COUNT(CASE WHEN errors_encountered IS NOT NULL AND errors_encountered != '[]' THEN 1 END) as decisions_with_errors
            FROM bot_decision_logs
            WHERE timestamp >= NOW() - INTERVAL '%s days'
            {bot_filter}
        """, params)

        row = cursor.fetchone()
        stats = {
            'total_decisions': row[0] or 0,
            'total_sessions': row[1] or 0,
            'entry_decisions': row[2] or 0,
            'exit_decisions': row[3] or 0,
            'skip_decisions': row[4] or 0,
            'profitable_trades': row[5] or 0,
            'losing_trades': row[6] or 0,
            'closed_trades': row[7] or 0,
            'avg_pnl': float(row[8]) if row[8] else 0,
            'total_pnl': float(row[9]) if row[9] else 0,
            'avg_slippage_pct': float(row[10]) if row[10] else 0,
            'avg_claude_response_ms': float(row[11]) if row[11] else 0,
            'decisions_with_errors': row[12] or 0
        }

        # Win rate
        if stats['closed_trades'] > 0:
            stats['win_rate'] = (stats['profitable_trades'] / stats['closed_trades']) * 100
        else:
            stats['win_rate'] = 0

        # By bot breakdown
        cursor.execute(f"""
            SELECT
                bot_name,
                COUNT(*) as count,
                COUNT(CASE WHEN actual_pnl > 0 THEN 1 END) as wins,
                SUM(actual_pnl) FILTER (WHERE actual_pnl IS NOT NULL) as total_pnl
            FROM bot_decision_logs
            WHERE timestamp >= NOW() - INTERVAL '%s days'
            GROUP BY bot_name
            ORDER BY count DESC
        """, [days])

        by_bot = []
        for row in cursor.fetchall():
            by_bot.append({
                'bot_name': row[0],
                'decisions': row[1],
                'wins': row[2] or 0,
                'total_pnl': float(row[3]) if row[3] else 0
            })

        cursor.close()
        conn.close()

        return {
            'success': True,
            'data': {
                'stats': stats,
                'by_bot': by_bot,
                'days_analyzed': days,
                'bot_filter': bot
            }
        }

    except Exception as e:
        cursor.close()
        conn.close()
        logger.error(f"Error getting bot decision stats: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/bot-decisions/export")
async def export_bot_decisions(
    bot: Optional[str] = None,
    format: str = Query("csv", regex="^(csv|json|excel)$"),
    days: int = 30,
    include_claude: bool = False  # Include full Claude prompts/responses (large!)
):
    """
    Export bot decisions to CSV, JSON, or Excel.

    Args:
        bot: Filter by bot name
        format: csv, json, or excel
        days: Number of days to export
        include_claude: Include full Claude prompts/responses (can be very large)
    """
    conn = get_connection()
    cursor = conn.cursor()

    try:
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables
                WHERE table_name = 'bot_decision_logs'
            )
        """)
        if not cursor.fetchone()[0]:
            raise HTTPException(status_code=400, detail="Table not created yet")

        bot_filter = "AND bot_name = %s" if bot else ""
        params = [days]
        if bot:
            params.append(bot.upper())

        # Select fields based on include_claude
        if include_claude:
            fields = "*"
        else:
            fields = """
                decision_id, bot_name, session_id, timestamp,
                decision_type, action, symbol, strategy,
                strike, expiration, option_type, contracts,
                spot_price, vix, net_gex, gex_regime,
                entry_reasoning, strike_reasoning,
                psychology_pattern, position_size_dollars,
                passed_all_checks, blocked_reason,
                order_submitted_at, order_filled_at,
                actual_fill_price, slippage_pct,
                actual_pnl, exit_triggered_by, outcome_notes
            """

        cursor.execute(f"""
            SELECT {fields}
            FROM bot_decision_logs
            WHERE timestamp >= NOW() - INTERVAL '%s days'
            {bot_filter}
            ORDER BY timestamp DESC
        """, params)

        columns = [desc[0] for desc in cursor.description]
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        if format == "json":
            data = []
            for row in rows:
                d = dict(zip(columns, row))
                # Convert datetime objects
                for key in d:
                    if isinstance(d[key], datetime):
                        d[key] = d[key].isoformat()
                data.append(d)
            return data

        elif format == "csv":
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(columns)

            for row in rows:
                # Convert datetime objects to strings
                processed_row = []
                for val in row:
                    if isinstance(val, datetime):
                        processed_row.append(val.isoformat())
                    elif isinstance(val, (dict, list)):
                        processed_row.append(json.dumps(val))
                    else:
                        processed_row.append(val)
                writer.writerow(processed_row)

            output.seek(0)
            filename = f"bot_decisions_{bot or 'all'}_{datetime.now().strftime('%Y%m%d')}.csv"

            return StreamingResponse(
                io.StringIO(output.getvalue()),
                media_type="text/csv",
                headers={"Content-Disposition": f"attachment; filename={filename}"}
            )

        elif format == "excel":
            # For Excel, we need openpyxl
            try:
                from openpyxl import Workbook
                from openpyxl.utils.dataframe import dataframe_to_rows
                import pandas as pd

                data = []
                for row in rows:
                    d = dict(zip(columns, row))
                    for key in d:
                        if isinstance(d[key], datetime):
                            d[key] = d[key].isoformat()
                        elif isinstance(d[key], (dict, list)):
                            d[key] = json.dumps(d[key])
                    data.append(d)

                df = pd.DataFrame(data)

                output = io.BytesIO()
                with pd.ExcelWriter(output, engine='openpyxl') as writer:
                    df.to_excel(writer, sheet_name='Bot Decisions', index=False)

                output.seek(0)
                filename = f"bot_decisions_{bot or 'all'}_{datetime.now().strftime('%Y%m%d')}.xlsx"

                return StreamingResponse(
                    output,
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={filename}"}
                )

            except ImportError:
                raise HTTPException(status_code=400, detail="Excel export requires openpyxl and pandas")

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error exporting bot decisions: {e}")
        raise HTTPException(status_code=500, detail=str(e))
